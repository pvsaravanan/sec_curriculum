import openpyxl
import pdfplumber
import json
import os
import re

excel_path = '/Users/saravanan.pv/Downloads/SEC - Learners -R19 & R24 - UG Curriculum - February 2025.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)

categories_standard = ['HS', 'BS', 'ES', 'PC', 'PE', 'OE', 'OL', 'EEC', 'MC', 'TOTAL']

def clean_val(v):
    if v is None:
        return ''
    if isinstance(v, float):
        if v.is_integer():
            return int(v)
        return round(v, 2)
    s = str(v).strip()
    return s

def num_val(v, default=0):
    if v is None:
        return default
    try:
        val = float(v)
        if val.is_integer():
            return int(val)
        return round(val, 2)
    except (ValueError, TypeError):
        return default

def parse_metadata(sheet, header_row, sheet_name):
    dept = ''
    degree = ''
    reg = ''
    title = ''
    batches = []
    
    for r in range(1, header_row):
        for c in range(1, 12):
            val = clean_val(sheet.cell(r, c).value)
            if not val:
                continue
            uval = str(val).upper()
            if 'DEPARTMENT OF' in uval:
                dept = val
            elif any(d in uval for d in ['B.E.', 'B.TECH', 'B.E ']):
                degree = val
            elif 'REGULATION' in uval:
                reg = val
            elif 'FLEXIBLE CHOICE BASED' in uval:
                title = val
            elif 'BATCH' in uval:
                if val not in batches:
                    batches.append(val)
                    
    # Clean up department and degree defaults if missing
    if not dept:
        name_clean = sheet_name.replace('R2019-', '').replace('R2024-', '').replace('2023-', '')
        dept_map = {
            'CIVIL': 'Department of Civil Engineering',
            'BME': 'Department of Biomedical Engineering',
            'MED': 'Department of Medical Electronics',
            'ECE': 'Department of Electronics and Communication Engineering',
            'EIE': 'Department of Electronics and Instrumentation Engineering',
            'EEE': 'Department of Electrical and Electronics Engineering',
            'CHEM': 'Department of Chemical Engineering',
            'AGRI': 'Department of Agricultural Engineering',
            'MECH': 'Department of Mechanical Engineering',
            'CSE': 'Department of Computer Science and Engineering',
            'IT': 'Department of Information Technology',
            'AI&DS': 'Department of Artificial Intelligence and Data Science',
            'AI&ML': 'Department of Artificial Intelligence and Machine Learning',
            'CSE(CS)': 'Department of Computer Science and Engineering (Cyber Security)',
            'CSE(IoT)': 'Department of Computer Science and Engineering (IoT)'
        }
        dept = dept_map.get(name_clean, f'Department of {name_clean}')

    if not degree:
        name_clean = sheet_name.replace('R2019-', '').replace('R2024-', '').replace('2023-', '')
        tech_depts = ['IT', 'AI&DS', 'AI&ML', 'CHEM']
        deg_prefix = 'B.Tech.' if name_clean in tech_depts else 'B.E.'
        degree = f"{deg_prefix} {sheet_name.split('-')[-1]}"

    if not reg:
        if 'R2024' in sheet_name:
            reg = 'REGULATIONS 2024'
        elif 'R2019' in sheet_name or '2023' in sheet_name:
            reg = 'REGULATIONS 2019'
        else:
            reg = 'AUTONOMOUS REGULATIONS'

    return {
        'department': dept,
        'degree': degree,
        'regulation': reg,
        'title': title or 'Flexible Choice Based Credit System (FCBCS)',
        'batches': batches
    }

def extract_pdf_online_courses(path):
    courses = []
    header_words_set = {'course', 'code', 'name', 'duration', 'provider', 'credits', 'equivalent', 'courses', 'certification', 'title', 'remarks'}
    with pdfplumber.open(path) as pdf:
        for page_idx, page in enumerate(pdf.pages):
            page_num = page_idx + 1
            words = page.extract_words()
            
            oc_words = [w for w in words if w['x0'] < 100 and re.match(r'^OC\d{3}$', w['text'])]
            oc_words.sort(key=lambda w: w['top'])
            
            for i, oc in enumerate(oc_words):
                top_b = oc_words[i-1]['top'] + 10 if i > 0 else (250 if page_num == 1 else (75 if page_num == 17 else 30))
                bottom_b = oc_words[i+1]['top'] - 8 if i+1 < len(oc_words) else page.height
                row_words = [w for w in words if top_b <= w['top'] < bottom_b and w['text'].lower() not in header_words_set]
                code = oc['text']
                
                if page_num <= 16:
                    name_words = [w['text'] for w in row_words if 100 <= w['x0'] < 265]
                    dur_words  = [w['text'] for w in row_words if 265 <= w['x0'] < 335]
                    prov_words = [w['text'] for w in row_words if 335 <= w['x0'] < 395]
                    cred_words = [w['text'] for w in row_words if 395 <= w['x0'] < 450]
                    equiv_words= [w['text'] for w in row_words if w['x0'] >= 450]
                else:
                    name_words = [w['text'] for w in row_words if 100 <= w['x0'] < 470]
                    dur_words  = [w['text'] for w in row_words if 470 <= w['x0'] < 535]
                    prov_words = [w['text'] for w in row_words if 535 <= w['x0'] < 640]
                    cred_words = [w['text'] for w in row_words if 640 <= w['x0'] < 690]
                    equiv_words= [w['text'] for w in row_words if w['x0'] >= 690]
                    
                title = ' '.join(name_words).strip()
                duration = ' '.join(dur_words).strip()
                provider = ' '.join(prov_words).strip()
                credits_str = ' '.join(cred_words).strip()
                equiv = ' '.join(equiv_words).strip()
                
                if not provider:
                    if page_num <= 10:
                        provider = 'NPTEL'
                    elif page_num <= 16:
                        provider = 'Coursera'
                
                try:
                    cred_num = int(credits_str)
                except ValueError:
                    if '12' in duration or '16' in duration or '24' in duration:
                        cred_num = 3
                    elif '8' in duration or '9' in duration or '10' in duration:
                        cred_num = 2
                    else:
                        cred_num = 1
                        
                remarks = f'{provider} • {duration}'
                if equiv and equiv != '-':
                    remarks += f' (Equiv: {equiv})'
                    
                courses.append({
                    'id': f'OnlineCourses_{code}',
                    'sno': len(courses) + 1,
                    'code_2024': code,
                    'code_2019': code,
                    'title': title,
                    'category': 'OL',
                    'theory_credits': cred_num,
                    'practical_credits': 0,
                    'total_credits': cred_num,
                    'prerequisite': '-',
                    'cgpa_type': 'CGPA',
                    'course_type': 'Online Course',
                    'remarks': remarks,
                    'provider': provider,
                    'duration': duration,
                    'equivalent_courses': equiv
                })
    return courses

def extract_batch_requirements(sheet, header_row, sheet_name, courses):
    offered_by_cat = {}
    for c in courses:
        cat = c.get('category', 'PC').split(' ')[0]
        offered_by_cat[cat] = offered_by_cat.get(cat, 0) + c.get('total_credits', 0)

    def attach_offered(req_list):
        for item in req_list:
            cat = item['category']
            if cat == 'TOTAL':
                item['offered_credits'] = sum(offered_by_cat.values())
            else:
                item['offered_credits'] = offered_by_cat.get(cat, 0)
        return req_list

    if sheet_name == 'R2024-BME':
        return {'Batch: 2024- 2028': attach_offered([
            {'category': 'HS', 'regular_credits': 10, 'lateral_credits': 2},
            {'category': 'BS', 'regular_credits': 23, 'lateral_credits': 6},
            {'category': 'ES', 'regular_credits': 30, 'lateral_credits': 16},
            {'category': 'PC', 'regular_credits': 60, 'lateral_credits': 60},
            {'category': 'PE', 'regular_credits': 18, 'lateral_credits': 18},
            {'category': 'OE', 'regular_credits': 4, 'lateral_credits': 4},
            {'category': 'OL', 'regular_credits': 3, 'lateral_credits': 3},
            {'category': 'EEC', 'regular_credits': 18, 'lateral_credits': 18},
            {'category': 'MC', 'regular_credits': 3, 'lateral_credits': 3},
            {'category': 'TOTAL', 'regular_credits': 169, 'lateral_credits': 130}
        ])}

    if sheet_name == 'R2024-CSE':
        return {'Batch: 2024- 2028': attach_offered([
            {'category': 'HS', 'regular_credits': 13, 'lateral_credits': 3},
            {'category': 'BS', 'regular_credits': 25, 'lateral_credits': 8},
            {'category': 'ES', 'regular_credits': 25, 'lateral_credits': 12},
            {'category': 'PC', 'regular_credits': 56, 'lateral_credits': 56},
            {'category': 'PE', 'regular_credits': 16, 'lateral_credits': 16},
            {'category': 'OE', 'regular_credits': 8, 'lateral_credits': 8},
            {'category': 'OL', 'regular_credits': 4, 'lateral_credits': 4},
            {'category': 'EEC', 'regular_credits': 16, 'lateral_credits': 16},
            {'category': 'MC', 'regular_credits': 3, 'lateral_credits': 3},
            {'category': 'TOTAL', 'regular_credits': 166, 'lateral_credits': 126}
        ])}

    if sheet_name == 'R2024-AI&DS':
        return {'Batch: 2024- 2028': attach_offered([
            {'category': 'HS', 'regular_credits': 14, 'lateral_credits': 3},
            {'category': 'BS', 'regular_credits': 23, 'lateral_credits': 8},
            {'category': 'ES', 'regular_credits': 28, 'lateral_credits': 12},
            {'category': 'PC', 'regular_credits': 56, 'lateral_credits': 56},
            {'category': 'PE', 'regular_credits': 16, 'lateral_credits': 16},
            {'category': 'OE', 'regular_credits': 8, 'lateral_credits': 8},
            {'category': 'OL', 'regular_credits': 4, 'lateral_credits': 4},
            {'category': 'EEC', 'regular_credits': 16, 'lateral_credits': 16},
            {'category': 'MC', 'regular_credits': 3, 'lateral_credits': 3},
            {'category': 'TOTAL', 'regular_credits': 168, 'lateral_credits': 126}
        ])}

    if sheet_name == 'R2024-AI&ML':
        return {'Batch: 2024- 2028': attach_offered([
            {'category': 'HS', 'regular_credits': 14, 'lateral_credits': 3},
            {'category': 'BS', 'regular_credits': 27, 'lateral_credits': 8},
            {'category': 'ES', 'regular_credits': 28, 'lateral_credits': 12},
            {'category': 'PC', 'regular_credits': 51, 'lateral_credits': 51},
            {'category': 'PE', 'regular_credits': 16, 'lateral_credits': 16},
            {'category': 'OE', 'regular_credits': 8, 'lateral_credits': 8},
            {'category': 'OL', 'regular_credits': 4, 'lateral_credits': 4},
            {'category': 'EEC', 'regular_credits': 16, 'lateral_credits': 16},
            {'category': 'MC', 'regular_credits': 3, 'lateral_credits': 3},
            {'category': 'TOTAL', 'regular_credits': 167, 'lateral_credits': 121}
        ])}

    if sheet_name == 'R2024-CSE(CS)':
        return {'Batch: 2024- 2028': attach_offered([
            {'category': 'HS', 'regular_credits': 14, 'lateral_credits': 3},
            {'category': 'BS', 'regular_credits': 25, 'lateral_credits': 8},
            {'category': 'ES', 'regular_credits': 28, 'lateral_credits': 12},
            {'category': 'PC', 'regular_credits': 56, 'lateral_credits': 56},
            {'category': 'PE', 'regular_credits': 16, 'lateral_credits': 16},
            {'category': 'OE', 'regular_credits': 8, 'lateral_credits': 8},
            {'category': 'OL', 'regular_credits': 4, 'lateral_credits': 4},
            {'category': 'EEC', 'regular_credits': 16, 'lateral_credits': 16},
            {'category': 'MC', 'regular_credits': 3, 'lateral_credits': 3},
            {'category': 'TOTAL', 'regular_credits': 170, 'lateral_credits': 126}
        ])}

    if sheet_name == 'R2024-IT':
        return {'Batch: 2024- 2028': attach_offered([
            {'category': 'HS', 'regular_credits': 13, 'lateral_credits': 3},
            {'category': 'BS', 'regular_credits': 25, 'lateral_credits': 8},
            {'category': 'ES', 'regular_credits': 25, 'lateral_credits': 12},
            {'category': 'PC', 'regular_credits': 56, 'lateral_credits': 56},
            {'category': 'PE', 'regular_credits': 16, 'lateral_credits': 16},
            {'category': 'OE', 'regular_credits': 8, 'lateral_credits': 8},
            {'category': 'OL', 'regular_credits': 4, 'lateral_credits': 4},
            {'category': 'EEC', 'regular_credits': 16, 'lateral_credits': 16},
            {'category': 'MC', 'regular_credits': 3, 'lateral_credits': 3},
            {'category': 'TOTAL', 'regular_credits': 166, 'lateral_credits': 126}
        ])}

    if sheet_name == 'R2019-MECH':
        return {
            'Batch: 2021- 2025': attach_offered([
                {'category': 'HS', 'regular_credits': 11, 'lateral_credits': 3},
                {'category': 'BS', 'regular_credits': 24, 'lateral_credits': 8},
                {'category': 'ES', 'regular_credits': 26, 'lateral_credits': 12},
                {'category': 'PC', 'regular_credits': 54, 'lateral_credits': 54},
                {'category': 'PE', 'regular_credits': 21, 'lateral_credits': 21},
                {'category': 'OE', 'regular_credits': 8, 'lateral_credits': 8},
                {'category': 'OL', 'regular_credits': 4, 'lateral_credits': 4},
                {'category': 'EEC', 'regular_credits': 16, 'lateral_credits': 16},
                {'category': 'MC', 'regular_credits': 3, 'lateral_credits': 3},
                {'category': 'TOTAL', 'regular_credits': 167, 'lateral_credits': 129}
            ]),
            'Batch: 2022- 2026': attach_offered([
                {'category': 'HS', 'regular_credits': 13, 'lateral_credits': 3},
                {'category': 'BS', 'regular_credits': 24, 'lateral_credits': 8},
                {'category': 'ES', 'regular_credits': 26, 'lateral_credits': 12},
                {'category': 'PC', 'regular_credits': 54, 'lateral_credits': 54},
                {'category': 'PE', 'regular_credits': 21, 'lateral_credits': 21},
                {'category': 'OE', 'regular_credits': 8, 'lateral_credits': 8},
                {'category': 'OL', 'regular_credits': 4, 'lateral_credits': 4},
                {'category': 'EEC', 'regular_credits': 16, 'lateral_credits': 16},
                {'category': 'MC', 'regular_credits': 3, 'lateral_credits': 3},
                {'category': 'TOTAL', 'regular_credits': 169, 'lateral_credits': 129}
            ])
        }

    if sheet_name == 'R2024-MECH':
        return {'Batch: 2024- 2028': attach_offered([
            {'category': 'HS', 'regular_credits': 10, 'lateral_credits': 4},
            {'category': 'BS', 'regular_credits': 25, 'lateral_credits': 6},
            {'category': 'ES', 'regular_credits': 29, 'lateral_credits': 17},
            {'category': 'PC', 'regular_credits': 55, 'lateral_credits': 55},
            {'category': 'PE', 'regular_credits': 18, 'lateral_credits': 18},
            {'category': 'OE', 'regular_credits': 0, 'lateral_credits': 0},
            {'category': 'OL', 'regular_credits': 0, 'lateral_credits': 0},
            {'category': 'EEC', 'regular_credits': 16, 'lateral_credits': 16},
            {'category': 'MC', 'regular_credits': 3, 'lateral_credits': 3},
            {'category': 'TOTAL', 'regular_credits': 156, 'lateral_credits': 119}
        ])}

    b1_title = 'Batch: 2021- 2025'
    b2_title = 'Batch: 2022- 2026, 2023-2027'
    for r in range(8, 11):
        c13 = sheet.cell(r, 13).value
        c16 = sheet.cell(r, 16).value
        if c13 and any(y in str(c13) for y in ['2019', '2021', '2022', '2023', '2024', '2025', '2026', '2027', '2028']):
            b1_title = str(c13).strip()
            if not b1_title.upper().startswith('BATCH'):
                b1_title = f"BATCH: {b1_title}"
        if c16 and any(y in str(c16) for y in ['2019', '2021', '2022', '2023', '2024', '2025', '2026', '2027', '2028']):
            b2_title = str(c16).strip()
            if not b2_title.upper().startswith('BATCH'):
                b2_title = f"BATCH: {b2_title}"
            
    cat_rows = {}
    for r in range(9, 25):
        val = str(sheet.cell(r, 13).value or '').strip().upper()
        if val in categories_standard:
            cat_rows[val] = r
            
    has_b2 = False
    if 'HS' in cat_rows:
        has_b2 = (sheet.cell(cat_rows['HS'], 17).value is not None or sheet.cell(cat_rows['HS'], 18).value is not None)
        
    b1_res = []
    b2_res = []
    
    is_aiml_2019 = (sheet_name == 'R2019-AI&ML')
    
    for cat in ['HS', 'BS', 'ES', 'PC', 'PE', 'OE', 'EEC', 'MC', 'TOTAL']:
        r = cat_rows.get(cat)
        if r:
            reg1 = sheet.cell(r, 14).value if sheet.cell(r, 14).value is not None else 0
            le1  = sheet.cell(r, 15).value if sheet.cell(r, 15).value is not None else 0
            if has_b2:
                if is_aiml_2019:
                    reg2 = sheet.cell(r, 18).value if sheet.cell(r, 18).value is not None else reg1
                    le2  = le1
                else:
                    reg2 = sheet.cell(r, 17).value if sheet.cell(r, 17).value is not None else reg1
                    le2  = sheet.cell(r, 18).value if sheet.cell(r, 18).value is not None else le1

            if cat == 'OE':
                reg1_v = int(reg1)
                le1_v = int(le1)
                oe1_r = 8 if reg1_v == 12 else (max(0, reg1_v - 2) if reg1_v == 6 else (4 if reg1_v == 7 else reg1_v))
                ol1_r = reg1_v - oe1_r
                oe1_l = 8 if le1_v == 12 else (max(0, le1_v - 2) if le1_v == 6 else (4 if le1_v == 7 else le1_v))
                ol1_l = le1_v - oe1_l
                
                b1_res.append({'category': 'OE', 'regular_credits': oe1_r, 'lateral_credits': oe1_l})
                b1_res.append({'category': 'OL', 'regular_credits': ol1_r, 'lateral_credits': ol1_l})
                
                if has_b2:
                    reg2_v = int(reg2)
                    le2_v = int(le2)
                    oe2_r = 8 if reg2_v == 12 else (max(0, reg2_v - 2) if reg2_v == 6 else (4 if reg2_v == 7 else reg2_v))
                    ol2_r = reg2_v - oe2_r
                    oe2_l = 8 if le2_v == 12 else (max(0, le2_v - 2) if le2_v == 6 else (4 if le2_v == 7 else le2_v))
                    ol2_l = le2_v - oe2_l
                    b2_res.append({'category': 'OE', 'regular_credits': oe2_r, 'lateral_credits': oe2_l})
                    b2_res.append({'category': 'OL', 'regular_credits': ol2_r, 'lateral_credits': ol2_l})
                continue

            b1_res.append({'category': cat, 'regular_credits': int(reg1), 'lateral_credits': int(le1)})
            if has_b2:
                b2_res.append({'category': cat, 'regular_credits': int(reg2), 'lateral_credits': int(le2)})
        else:
            b1_res.append({'category': cat, 'regular_credits': 0, 'lateral_credits': 0})
            if has_b2:
                b2_res.append({'category': cat, 'regular_credits': 0, 'lateral_credits': 0})
                
    result = {b1_title: attach_offered(b1_res)}
    if has_b2:
        result[b2_title] = attach_offered(b2_res)
    return result

sheets_data = []

for sheet_name in wb.sheetnames:
    sheet = wb[sheet_name]
    
    if sheet_name == 'Minors':
        sheet_type = 'minors'
        header_row = 8
        meta = {
            'department': 'Inter-Departmental Specializations',
            'degree': 'Specialization (Minor Degree)',
            'regulation': 'Regulations 2019 & R2024',
            'title': 'List of Courses Offered as Specialization (Minor)',
            'batches': []
        }
        courses = []
        for r in range(header_row + 1, sheet.max_row + 1):
            sno = clean_val(sheet.cell(r, 1).value)
            offered_by = clean_val(sheet.cell(r, 2).value)
            specialization = clean_val(sheet.cell(r, 3).value)
            code_19 = clean_val(sheet.cell(r, 4).value)
            title = clean_val(sheet.cell(r, 5).value)
            if not title or title.lower() == 'course title':
                continue
            l_hrs = num_val(sheet.cell(r, 6).value)
            t_hrs = num_val(sheet.cell(r, 7).value)
            p_hrs = num_val(sheet.cell(r, 8).value)
            credits = num_val(sheet.cell(r, 9).value)
            prereq = clean_val(sheet.cell(r, 10).value)
            cgpa_type = clean_val(sheet.cell(r, 11).value)
            course_type = clean_val(sheet.cell(r, 12).value)
            courses.append({
                'id': f"{sheet_name}_c{len(courses)}",
                'sno': sno,
                'offered_by': offered_by,
                'specialization': specialization,
                'code_2019': code_19,
                'code_2024': '',
                'title': title,
                'l_hrs': l_hrs,
                't_hrs': t_hrs,
                'p_hrs': p_hrs,
                'theory_credits': l_hrs,
                'practical_credits': p_hrs,
                'total_credits': credits,
                'prerequisite': prereq or '-',
                'cgpa_type': cgpa_type or 'CGPA',
                'course_type': course_type or 'Theory',
                'category': 'Minor',
                'remarks': f'Specialization in {specialization} (Offered by {offered_by})'
            })
        req_list = [
            {'category': 'Minor Specialization', 'regular_credits': 18, 'lateral_credits': 18, 'offered_credits': sum(c['total_credits'] for c in courses)},
            {'category': 'TOTAL', 'regular_credits': 18, 'lateral_credits': 18, 'offered_credits': sum(c['total_credits'] for c in courses)}
        ]
        sheets_data.append({
            'sheet_name': sheet_name,
            'sheet_type': sheet_type,
            'metadata': meta,
            'batches': {'All Batches': req_list},
            'credit_requirements': req_list,
            'courses': courses
        })

    elif sheet_name == 'Open Electives':
        sheet_type = 'open_electives'
        header_row = 7
        meta = {
            'department': 'Inter-Departmental Offerings',
            'degree': 'Open Electives (OE)',
            'regulation': 'Regulations 2019 & R2024',
            'title': 'List of Open Electives for B.E. / B.Tech.',
            'batches': []
        }
        courses = []
        for r in range(header_row + 1, sheet.max_row + 1):
            sno = clean_val(sheet.cell(r, 1).value)
            offered_by = clean_val(sheet.cell(r, 2).value)
            code_24 = clean_val(sheet.cell(r, 3).value)
            code_19 = clean_val(sheet.cell(r, 4).value)
            title = clean_val(sheet.cell(r, 5).value)
            if not title or title.lower() == 'course title':
                continue
            prereq = clean_val(sheet.cell(r, 6).value)
            th_credits = num_val(sheet.cell(r, 7).value)
            pr_credits = num_val(sheet.cell(r, 8).value)
            tot_credits = num_val(sheet.cell(r, 9).value)
            if tot_credits == 0 and (th_credits > 0 or pr_credits > 0):
                tot_credits = th_credits + pr_credits
            courses.append({
                'id': f"{sheet_name}_c{len(courses)}",
                'sno': sno,
                'offered_by': offered_by,
                'code_2024': code_24,
                'code_2019': code_19,
                'title': title,
                'prerequisite': prereq or '-',
                'theory_credits': th_credits,
                'practical_credits': pr_credits,
                'total_credits': tot_credits,
                'category': 'OE',
                'cgpa_type': 'CGPA',
                'course_type': 'Theory',
                'remarks': f'Offered by Department of {offered_by}'
            })
        oe_req_list = [
            {'category': 'OE', 'regular_credits': 8, 'lateral_credits': 8, 'offered_credits': sum(c['total_credits'] for c in courses)},
            {'category': 'TOTAL', 'regular_credits': 8, 'lateral_credits': 8, 'offered_credits': sum(c['total_credits'] for c in courses)}
        ]
        sheets_data.append({
            'sheet_name': sheet_name,
            'sheet_type': sheet_type,
            'metadata': meta,
            'batches': {'All Batches': oe_req_list},
            'credit_requirements': oe_req_list,
            'courses': courses
        })

    elif sheet_name == 'Online Courses':
        sheet_type = 'online_courses'
        header_row = 1
        meta = {
            'department': 'Academic Committee (NPTEL / Coursera / Oracle / RedHat / Microsoft / Cisco)',
            'degree': 'Online Courses (OL) - 4 Credits',
            'regulation': 'Regulations 2019 & R2024 (2024 EVEN)',
            'title': 'Approved List of Online Courses for CSE/IT/AI&DS/AI&ML/CSE(IoT)/CSE(CS)',
            'batches': []
        }
        pdf_courses_path = '/Users/saravanan.pv/Downloads/Online Courses - Academic Committee - AI_CS_IT @024 EVEN.pdf'
        courses = extract_pdf_online_courses(pdf_courses_path)
        ol_req_list = [
            {'category': 'OL', 'regular_credits': 4, 'lateral_credits': 4, 'offered_credits': sum(c['total_credits'] for c in courses)},
            {'category': 'TOTAL', 'regular_credits': 4, 'lateral_credits': 4, 'offered_credits': sum(c['total_credits'] for c in courses)}
        ]
        sheets_data.append({
            'sheet_name': sheet_name,
            'sheet_type': sheet_type,
            'metadata': meta,
            'batches': {'All Batches': ol_req_list},
            'credit_requirements': ol_req_list,
            'courses': courses
        })

    else:
        # Standard curriculum sheet
        sheet_type = 'curriculum'
        header_row = 10
        for r in range(1, 15):
            vals = [str(sheet.cell(r, c).value or '').strip() for c in range(1, 12)]
            if any('Course Code' in v or 'Course Title' in v for v in vals):
                header_row = r
                break
        
        meta = parse_metadata(sheet, header_row, sheet_name)
        
        courses = []
        for r in range(header_row + 1, sheet.max_row + 1):
            sno = clean_val(sheet.cell(r, 1).value)
            remarks = clean_val(sheet.cell(r, 2).value)
            cat = clean_val(sheet.cell(r, 3).value).upper()
            code_24 = clean_val(sheet.cell(r, 4).value)
            code_19 = clean_val(sheet.cell(r, 5).value)
            title = clean_val(sheet.cell(r, 6).value)
            if not title or title.lower() in ['course title', 'title', 'none', 's.no']:
                continue
            prereq = clean_val(sheet.cell(r, 7).value)
            th_credits = num_val(sheet.cell(r, 8).value)
            pr_credits = num_val(sheet.cell(r, 9).value)
            tot_credits = num_val(sheet.cell(r, 10).value)
            if tot_credits == 0 and (th_credits > 0 or pr_credits > 0):
                tot_credits = th_credits + pr_credits
            cgpa_type = clean_val(sheet.cell(r, 11).value)
            course_type = clean_val(sheet.cell(r, 12).value)
            
            if not cat:
                cat = 'PC'
            
            courses.append({
                'id': f"{sheet_name}_c{len(courses)}",
                'sno': sno,
                'remarks': remarks,
                'category': cat,
                'code_2024': code_24,
                'code_2019': code_19,
                'title': title,
                'prerequisite': prereq or '-',
                'theory_credits': th_credits,
                'practical_credits': pr_credits,
                'total_credits': tot_credits,
                'cgpa_type': cgpa_type or 'CGPA',
                'course_type': course_type or 'Theory'
            })
            
        batch_dict = extract_batch_requirements(sheet, header_row, sheet_name, courses)
        # Default to the latest active student batch (e.g. 2022-2026, 2023-2027 / 2024-2028)
        default_batch = list(batch_dict.keys())[-1]
        credit_reqs = batch_dict[default_batch]
            
        sheets_data.append({
            'sheet_name': sheet_name,
            'sheet_type': sheet_type,
            'metadata': meta,
            'active_batch': default_batch,
            'batches': batch_dict,
            'credit_requirements': credit_reqs,
            'courses': courses
        })

output_path = '/Users/saravanan.pv/Desktop/sec-curriculum-dashboard/curriculum_data.json'
with open(output_path, 'w', encoding='utf-8') as f:
    json.dump(sheets_data, f, ensure_ascii=False, indent=2)

print(f'Successfully re-extracted {len(sheets_data)} sheets to {output_path}')
print('Total courses across all sheets:', sum(len(s['courses']) for s in sheets_data))
